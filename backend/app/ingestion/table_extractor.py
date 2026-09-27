"""Extract CSV and XLSX files into page-scoped table text."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path


@dataclass
class TablePage:
    page_number: int
    title: str
    text: str


def extract_table_file(path: Path, extension: str) -> list[TablePage]:
    if extension == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return [_to_page(1, path.stem, rows)]

    if extension == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX support requires openpyxl. Install backend requirements.") from exc

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            pages: list[TablePage] = []
            for page_number, worksheet in enumerate(workbook.worksheets, start=1):
                rows = [[("" if value is None else str(value)) for value in row] for row in worksheet.iter_rows(values_only=True)]
                pages.append(_to_page(page_number, worksheet.title, rows))
            return pages
        finally:
            workbook.close()

    raise ValueError(f"Unsupported table extension '.{extension}'.")


def _to_page(page_number: int, title: str, rows: list[list[str]]) -> TablePage:
    cleaned_rows = [[cell.strip() for cell in row] for row in rows]
    while cleaned_rows and not any(cleaned_rows[0]):
        cleaned_rows.pop(0)
    while cleaned_rows and not any(cleaned_rows[-1]):
        cleaned_rows.pop()

    if not cleaned_rows:
        return TablePage(page_number=page_number, title=title, text=f"Table: {title}\n(empty)")

    width = max(len(row) for row in cleaned_rows)
    normalized = [row + [""] * (width - len(row)) for row in cleaned_rows]
    headers = normalized[0]
    body = normalized[1:]
    lines = [f"Table: {title}", "Headers: " + " | ".join(headers)]
    lines.extend(" | ".join(row) for row in body if any(row))
    return TablePage(page_number=page_number, title=title, text="\n".join(lines))
