"""Extract searchable text and tables from DOCX files."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class DocxPage:
    page_number: int
    text: str


def extract_docx(path: Path) -> list[DocxPage]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx. Install backend requirements.") from exc

    document = Document(path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"Table: {table_index}")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    text = "\n".join(lines).strip()
    return [DocxPage(page_number=1, text=text or "(empty document)")]
